#!/usr/bin/env python3
"""通知草稿（参照替换式）：以参照机构通知 docx 为底，清空正文后按 md 重建段落。

原理：不仿制样式，直接克隆参照文档中同类段落的 <w:p>（pPr 缩进行距、rPr 字体一并
继承），只替换文字；参照文档的红头区、文号、表格、图片随清空一并移除；页面设置、
节属性原生保留；页眉清空（避免带入对方红头），页脚保留。

捐助段落（donor）：标题（小标宋居中）、称谓（不缩进）、正文（首行缩进两字）、
一级标题（黑体）、二级标题（楷体）、落款（右对齐）六类。换参照 docx 时先用
--inspect 定位，再用 --title/--salute/... 指定段号。

md 约定：# 标题；第一个非标题段为称谓；## 一级；### 二级；普通段为正文；
末尾"机构＋日期"两行自动识别为落款；**粗体**生效；【待核…】黄色高亮；
文件头 HTML 注释块自动跳过。

用法：
  python notice_draft_docx.py 参照.docx --inspect
  python notice_draft_docx.py 参照.docx 输入.md 输出.docx \
      --title 4 --salute 6 --body 7 --h1 15 --h2 30 --sign 68 \
      --font-override 楷体_GB2312=楷体
段号缺省值对应《中国高等教育培训中心》研修班通知（2026-10 首个实战参照）。
"""
import argparse
import copy
import re

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def build_segments(text):
    """按【待核…】与**粗体**切分为 (文字, 粗体, 高亮) 列表。"""
    segs = []
    for a in re.split(r"(【待核[^】]*】)", text):
        if not a:
            continue
        hl = a.startswith("【待核")
        for b in re.split(r"(\*\*.+?\*\*)", a):
            if not b:
                continue
            bold = b.startswith("**")
            segs.append((b[2:-2] if bold else b, bold, hl))
    return segs


def make_run(template, text, bold, highlight, font_overrides):
    r = copy.deepcopy(template)
    for tag in ("w:drawing", "w:br", "w:cr", "w:tab", "w:t"):
        for e in r.findall(qn(tag)):
            r.remove(e)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    rpr = r.find(qn("w:rPr"))
    if rpr is None:
        rpr = OxmlElement("w:rPr")
        r.insert(0, rpr)
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None and font_overrides:
        for attr in ("w:eastAsia", "w:ascii", "w:hAnsi", "w:cs"):
            val = rfonts.get(qn(attr))
            if val in font_overrides:
                rfonts.set(qn(attr), font_overrides[val])
    if bold:
        b = OxmlElement("w:b")
        rfonts2 = rpr.find(qn("w:rFonts"))
        rfonts2.addnext(b) if rfonts2 is not None else rpr.insert(0, b)
    if highlight:
        h = OxmlElement("w:highlight")
        h.set(qn("w:val"), "yellow")
        rpr.append(h)
    return r


def para_from_donor(donor, text, font_overrides):
    p = copy.deepcopy(donor)
    runs = p.findall(qn("w:r"))
    if not runs:
        raise RuntimeError("参照段落无 run，不能作捐助段落，请换一段")
    template = copy.deepcopy(runs[0])
    for r in runs:
        p.remove(r)
    for link in p.findall(qn("w:hyperlink")):
        p.remove(link)
    for text_, bold, hl in build_segments(text):
        p.append(make_run(template, text_, bold, hl, font_overrides))
    return p


def blank_from_donor(donor):
    p = copy.deepcopy(donor)
    for r in p.findall(qn("w:r")):
        p.remove(r)
    return p


def read_md_content(md_path):
    lines = open(md_path, encoding="utf-8").read().splitlines()
    content, in_comment = [], False
    for ln in lines:
        s = ln.strip()
        if s.startswith("<!--"):
            in_comment = True
            continue
        if in_comment:
            if "-->" in s:
                in_comment = False
            continue
        if s:
            content.append(s)
    sign = []
    if len(content) >= 2 and re.match(r"^\d{4}年\d{1,2}月\d{1,2}日", content[-1]):
        sign = content[-2:]
        content = content[:-2]
    return content, sign


def inspect(ref_path):
    doc = Document(ref_path)
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip():
            print(i, "|", p.style.name, "|", p.text.strip()[:40])


def convert(args):
    doc = Document(args.ref)
    paras = doc.paragraphs
    idx = {"title": args.title, "salute": args.salute, "body": args.body,
           "h1": args.h1, "h2": args.h2, "sign": args.sign}
    donors = {k: paras[i]._p for k, i in idx.items()}
    for k, i in idx.items():
        label = donors[k].xpath("string(.)")[:24]
        print(f"捐助段落[{k}]={i}: {label}")

    body = doc.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    sect_pr = body.find(qn("w:sectPr"))

    def emit(p_elem):
        sect_pr.addprevious(p_elem)

    content, sign = read_md_content(args.md)
    first_para_seen = False
    for s in content:
        if s.startswith("# "):
            emit(para_from_donor(donors["title"], s[2:], args.font_overrides))
            emit(blank_from_donor(donors["salute"]))
        elif s.startswith("## "):
            emit(para_from_donor(donors["h1"], s[3:], args.font_overrides))
        elif s.startswith("### "):
            emit(para_from_donor(donors["h2"], s[4:], args.font_overrides))
        elif not first_para_seen:
            emit(para_from_donor(donors["salute"], s, args.font_overrides))
            first_para_seen = True
        else:
            emit(para_from_donor(donors["body"], s, args.font_overrides))

    if sign:
        emit(blank_from_donor(donors["salute"]))
        for s in sign:
            emit(para_from_donor(donors["sign"], s, args.font_overrides))

    for sec in doc.sections:
        for p in sec.header.paragraphs:
            for r in list(p.runs):
                r._element.getparent().remove(r._element)

    doc.save(args.out)
    print(f"已生成 {args.out}，段落数 {len(doc.paragraphs)}")


def main():
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("ref", help="参照 docx 路径")
    ap.add_argument("md", nargs="?", help="输入 md 路径（--inspect 时省略）")
    ap.add_argument("out", nargs="?", help="输出 docx 路径")
    ap.add_argument("--inspect", action="store_true", help="仅列出参照文档段落，用于定位捐助段落")
    ap.add_argument("--title", type=int, default=4, help="标题捐助段落号（0 起，下同）")
    ap.add_argument("--salute", type=int, default=6, help="称谓捐助段落号")
    ap.add_argument("--body", type=int, default=7, help="正文捐助段落号")
    ap.add_argument("--h1", type=int, default=15, help="一级标题捐助段落号")
    ap.add_argument("--h2", type=int, default=30, help="二级标题捐助段落号")
    ap.add_argument("--sign", type=int, default=68, help="落款捐助段落号")
    ap.add_argument("--font-override", action="append", default=[],
                    metavar="原名=新名", help="字体名替换，可多次，如 楷体_GB2312=楷体")
    args = ap.parse_args()
    if args.inspect:
        inspect(args.ref)
        return
    if not (args.md and args.out):
        ap.error("转换模式需要 md 与 out 两个参数")
    args.font_overrides = dict(kv.split("=", 1) for kv in args.font_override)
    convert(args)


if __name__ == "__main__":
    main()
