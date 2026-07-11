"""Header and footer setup with multi-section support."""
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def _add_page_number_field(p, cfg: dict):
    """Add a PAGE field to a paragraph with surrounding dashes."""
    hf_cfg = cfg.get('header_footer', {})
    font_size = Pt(hf_cfg.get('footer_font_size', 9))
    font_color = RGBColor.from_string(hf_cfg.get('footer_color', '999999'))

    run = p.add_run('— ')
    run.font.size = font_size
    run.font.color.rgb = font_color

    # PAGE field
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run1 = p.add_run()
    run1._element.append(fldChar1)

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' PAGE '
    run2 = p.add_run()
    run2._element.append(instrText)

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run3 = p.add_run()
    run3._element.append(fldChar2)

    run4 = p.add_run(' —')
    run4.font.size = font_size
    run4.font.color.rgb = font_color


def _set_page_number_start(section, start: int = 1):
    """Set page number to start from a specific value in a section."""
    sectPr = section._sectPr
    pgNumType = sectPr.find(qn('w:pgNumType'))
    if pgNumType is None:
        pgNumType = OxmlElement('w:pgNumType')
        sectPr.append(pgNumType)
    pgNumType.set(qn('w:start'), str(start))


def setup_header_footer(doc, meta, cfg: dict, has_cover: bool = False):
    """Set up page header and footer based on template config.

    Args:
        has_cover: If True, section 0 is the cover (no page number),
                   and page numbers start from section 1.
    """
    hf_cfg = cfg.get('header_footer', {})
    header_text = meta.header_text or meta.title

    if has_cover and len(doc.sections) > 1:
        # Section 0: cover page — clear header/footer
        cover_section = doc.sections[0]
        cover_section.header.is_linked_to_previous = False
        cover_section.footer.is_linked_to_previous = False
        # Ensure cover has empty header/footer
        _clear_header_footer(cover_section)

        # Section 1+: body — set header/footer with page number from 1
        body_section = doc.sections[1]
        _setup_section(body_section, header_text, hf_cfg, cfg)
        _set_page_number_start(body_section, start=1)
    else:
        # Single section document
        section = doc.sections[0]
        _setup_section(section, header_text, hf_cfg, cfg)


def _clear_header_footer(section):
    """Clear header and footer for a section (e.g. cover page)."""
    # Clear header
    header = section.header
    header.is_linked_to_previous = False
    for p in header.paragraphs:
        for run in p.runs:
            run.text = ''

    # Clear footer
    footer = section.footer
    footer.is_linked_to_previous = False
    for p in footer.paragraphs:
        for run in p.runs:
            run.text = ''


def _setup_section(section, header_text: str, hf_cfg: dict, cfg: dict):
    """Set up header and footer for a single section."""
    # Header
    if header_text and hf_cfg.get('header_visible', True):
        header = section.header
        header.is_linked_to_previous = False
        p = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = p.add_run(header_text)
        run.font.size = Pt(hf_cfg.get('header_font_size', 9))
        color_str = hf_cfg.get('header_color', '999999')
        run.font.color.rgb = RGBColor.from_string(color_str)

    # Footer with page number
    footer_type = hf_cfg.get('footer_type', 'page_number')
    if footer_type == 'page_number':
        footer = section.footer
        footer.is_linked_to_previous = False
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_field(p, cfg)
