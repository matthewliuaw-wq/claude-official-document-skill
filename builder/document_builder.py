"""Document builder - orchestrates all modules to produce a formatted .docx."""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from .styles import apply_body_style, apply_heading_styles, configure_page, set_east_asian_font
from .numbering import create_heading_numbering
from .tables import build_table
from .cover_page import add_cover_page, add_separator_line
from .header_footer import setup_header_footer
from .validator import validate_document, format_validation_report


def _convert_chinese_quotes(text: str) -> str:
    """Convert straight quotes to paired Chinese quotes.

    " → "" (alternating open/close)
    ' → '' (alternating open/close)
    Already-Chinese quotes are left untouched.
    """
    result = []
    d_open = False
    s_open = False
    for ch in text:
        if ch == '"':
            result.append('“' if not d_open else '”')
            d_open = not d_open
        elif ch == "'":
            result.append('‘' if not s_open else '’')
            s_open = not s_open
        else:
            result.append(ch)
    return ''.join(result)


def _maybe_convert_quotes(text: str, cfg: dict) -> str:
    """Apply text processing pipeline if the template config enables it."""
    if not text:
        return text
    tp = cfg.get('text_processing', {})
    if not tp:
        return text
    # Quote conversion
    if tp.get('convert_quotes'):
        text = _convert_chinese_quotes(text)
    # Ellipsis normalization
    if tp.get('normalize_ellipsis'):
        text = _normalize_ellipsis(text)
    # Common error fixes
    if tp.get('fix_common_errors'):
        text = _fix_common_errors(text)
    return text


def _normalize_ellipsis(text: str) -> str:
    """Normalize ellipsis to Chinese standard form.

    ...... or . . . . . . → …… (6 dots, centered)
    ... or . . .           → … (3 dots, centered)
    """
    import re
    # Replace 6+ consecutive dots with Chinese ellipsis
    text = re.sub(r'\.{6,}', '……', text)
    # Replace exactly 3-5 dots with 3-dot Chinese ellipsis
    text = re.sub(r'\.{3,5}', '…', text)
    # Replace spaced dots: . . . . . . → ……
    text = re.sub(r'(\.\s+){5}\.', '……', text)
    text = re.sub(r'(\.\s+){2}\.', '…', text)
    return text


def _fix_common_errors(text: str) -> str:
    """Fix common Chinese writing errors per publisher standards."""
    replacements = [
        ('其它', '其他'),           # "其它" is non-standard
    ]
    for old, new in replacements:
        text = text.replace(old, new)
    return text


def build_document(content, template_cfg):
    """Build a complete python-docx Document from content + template config."""
    doc = Document()

    # 1. Page setup
    configure_page(doc, template_cfg)

    # 2. Apply body and heading styles
    apply_body_style(doc, template_cfg)
    apply_heading_styles(doc, template_cfg)

    # 3. Create heading numbering
    num_id = create_heading_numbering(doc, template_cfg)

    # 4. Cover page
    has_cover = bool(content.cover)
    if has_cover:
        add_cover_page(doc, content.cover, template_cfg)
        add_separator_line(doc, template_cfg)
        # Insert section break: cover page has no page number, body starts from 1
        new_section = doc.add_section(WD_SECTION.NEW_PAGE)
        new_section.header.is_linked_to_previous = False
        new_section.footer.is_linked_to_previous = False
        # Re-apply page setup for new section (margins etc.)
        configure_page(doc, template_cfg)

    # 5. Process content blocks
    for block in content.content:
        btype = block.type

        if btype == 'heading':
            _add_heading(doc, block, template_cfg)

        elif btype == 'paragraph':
            _add_paragraph(doc, block, template_cfg)

        elif btype == 'table':
            if block.table:
                build_table(doc, block.table, template_cfg)

        elif btype == 'bullet_list':
            if block.bullet_list:
                _add_bullet_list(doc, block.bullet_list)

        elif btype == 'code_block':
            if block.code_block:
                _add_code_block(doc, block.code_block, template_cfg)

        elif btype == 'figure':
            if block.figure:
                _add_figure(doc, block.figure, template_cfg)

        elif btype == 'page_break':
            doc.add_page_break()

        elif btype == 'toc':
            _add_toc(doc)

        elif btype == 'attachment_label':
            _add_attachment_label(doc, block, template_cfg)

    # 6. Header & footer (handles multi-section for cover vs body)
    setup_header_footer(doc, content.meta, template_cfg, has_cover=has_cover)

    return doc


def _is_front_matter_heading(text: str) -> bool:
    """Check if H1 text is a front-matter or back-matter heading (not a chapter title).

    Front/back-matter headings (内容提要, 前言, 目录, 参考文献 etc.) use a smaller
    font size than chapter titles per publisher standards.
    """
    import re
    # Strip spaces used for visual spacing (e.g. 前　言, 目　录)
    stripped = text.replace('　', '').replace(' ', '')
    front_matter_patterns = [
        '内容提要', '前言', '序言', '序', '目录',
        '参考文献', '附录', '后记', '索引',
    ]
    if stripped in front_matter_patterns:
        return True
    return False


def _add_heading(doc, block, cfg):
    """Add a heading paragraph with East Asian font, per-level font support."""
    from docx.enum.text import WD_LINE_SPACING
    from docx.shared import RGBColor
    level = block.level or 1
    text = _maybe_convert_quotes(block.text or '', cfg)

    # Front-matter H1 headings (前言/目录/参考文献) start on a new page.
    # 内容提要 is the first H1 — it follows the cover page section break, no extra break needed.
    if level == 1 and _is_front_matter_heading(text):
        h1_count = sum(1 for p in doc.paragraphs if p.style.name == 'Heading 1')
        if h1_count > 0:  # not the first H1
            doc.add_page_break()

    heading = doc.add_heading(text, level=level)

    heading_cfg = cfg.get('fonts', {}).get('heading', {})
    east = heading_cfg.get('east_asia', 'SimHei')
    level_east = heading_cfg.get(f'h{level}_east_asia', east)
    latin = heading_cfg.get('latin', 'Times New Roman')

    # Determine run-level size override
    # H1 front-matter headings use a smaller size than chapter titles
    run_size = None
    if level == 1 and _is_front_matter_heading(text):
        run_size = heading_cfg.get('h1_front_matter_size')

    # Heading color — explicit black to avoid theme color inheritance
    heading_color_str = heading_cfg.get('color', '000000')
    heading_color = RGBColor.from_string(heading_color_str)

    for run in heading.runs:
        set_east_asian_font(run._element, latin, level_east, run_size, unify_fonts=True)
        run.font.color.rgb = heading_color

    # Per-level bold control on runs (must override python-docx default bold)
    bold_val = heading_cfg.get(f'h{level}_bold')
    if bold_val is not None:
        for run in heading.runs:
            run.font.bold = bold_val
            # Fix: python-docx adds empty <w:bCs/> which defaults to true.
            # Must explicitly set bCs to match b.
            rPr = run._element.find(qn('w:rPr'))
            if rPr is not None:
                for bcs in rPr.findall(qn('w:bCs')):
                    rPr.remove(bcs)
                bcs_el = rPr.makeelement(qn('w:bCs'), {qn('w:val'): str(bold_val).lower()})
                rPr.append(bcs_el)

    # Per-level alignment: clear indent for centered or left-aligned headings
    level_align = heading_cfg.get(f'h{level}_alignment')
    if level_align in ('center', 'left'):
        heading.paragraph_format.left_indent = Pt(0)
        heading.paragraph_format.first_line_indent = Pt(0)

    # GB/T 9704: title (level 1) line spacing override
    # Two lines: single line spacing; more than two lines: fixed 35pt
    title_ls = cfg.get('paragraph', {}).get('title_line_spacing')
    if title_ls and level == 1:
        pf = heading.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        if title_ls == 'single':
            pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        elif isinstance(title_ls, str) and title_ls.endswith('pt'):
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(float(title_ls.replace('pt', '')))


def _is_level3_paragraph(text: str) -> bool:
    """Check if text starts with a level-3 pattern: '1.' '2.' etc. (Arabic numeral + period)."""
    import re
    return bool(re.match(r'^\d+\.', text))


def _add_paragraph(doc, block, cfg):
    """Add a body paragraph with optional bold prefix and alignment override."""
    text = _maybe_convert_quotes(block.text or '', cfg)
    p = doc.add_paragraph()

    # Per-block alignment override (also clears first-line indent)
    if block.alignment:
        align_map = {
            'justified': WD_ALIGN_PARAGRAPH.JUSTIFY,
            'left': WD_ALIGN_PARAGRAPH.LEFT,
            'center': WD_ALIGN_PARAGRAPH.CENTER,
            'right': WD_ALIGN_PARAGRAPH.RIGHT,
        }
        p.alignment = align_map.get(block.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
        p.paragraph_format.first_line_indent = Pt(0)

    if block.bold_prefix:
        run = p.add_run(block.bold_prefix)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 0, 0)

    run = p.add_run(text)
    run.font.color.rgb = RGBColor(0, 0, 0)

    # GB/T 9704: level-3 paragraphs (1. 2. etc.) should be bold in official documents
    if _is_level3_paragraph(text):
        run.font.bold = True

    # Per-block space_before override
    if block.space_before:
        p.paragraph_format.space_before = Pt(block.space_before)

    return p


def _add_bullet_list(doc, list_data):
    """Add bullet or numbered list items."""
    style = list_data.style
    for item in list_data.items:
        if style == 'numbered':
            doc.add_paragraph(item, style='List Number')
        else:
            doc.add_paragraph(item, style='List Bullet')


def _add_attachment_label(doc, block, cfg):
    """Add an attachment label (e.g. '附件1') in heading font, left-aligned, no indent.

    GB/T 9704: attachment label in 黑体 三号 (16pt), top-left corner of attachment page.
    """
    text = block.text or ''
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Pt(0)

    run = p.add_run(text)
    run.font.bold = True

    # Use heading font (黑体) and size (h2 = 16pt for 公文)
    heading_cfg = cfg.get('fonts', {}).get('heading', {})
    east = heading_cfg.get('east_asia', 'SimHei')
    latin = heading_cfg.get('latin', 'Times New Roman')
    size = heading_cfg.get('h2_size', 16)
    set_east_asian_font(run._element, latin, east, size, unify_fonts=True)


def _add_toc(doc):
    """Insert a Table of Contents field."""
    from docx.shared import RGBColor
    p = doc.add_paragraph()

    # TOC field runs — all share the same formatting
    toc_size = Pt(10.5)   # 五号，与正文一致
    toc_color = RGBColor(0, 0, 0)  # 黑色

    run = p.add_run()
    run.font.size = toc_size
    run.font.color.rgb = toc_color

    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    run._element.append(fldChar1)

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = ' TOC \\o "1-3" \\h \\z \\u '
    run2 = p.add_run()
    run2.font.size = toc_size
    run2.font.color.rgb = toc_color
    run2._element.append(instrText)

    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'end')
    run3 = p.add_run()
    run3.font.size = toc_size
    run3.font.color.rgb = toc_color
    run3._element.append(fldChar2)


def _add_code_block(doc, code_data, cfg):
    """Add a code block with monospace font, light gray background.

    Publisher standard: code uses smaller font (9pt), no indent, left-aligned.
    """
    from docx.shared import RGBColor
    code = code_data.code or ''

    for line in code.split('\n'):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.first_line_indent = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

        # Light gray background shading
        pPr = p._element.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), 'F5F5F5')
        pPr.append(shd)

        run = p.add_run(line if line else ' ')  # avoid empty run
        run.font.size = Pt(9)
        run.font.name = 'Consolas'
        set_east_asian_font(run._element, 'Consolas', 'SimSun', 9)

    # Optional language label above (if we had added it before the code)


def _add_figure(doc, fig_data, cfg):
    """Add a figure with image or placeholder, plus caption below.

    If image_path is provided and file exists, inserts the image centered.
    Otherwise, inserts a placeholder box.
    Caption format: centered, '图X-Y 图名', 9pt SimSun.
    """
    import os
    from docx.shared import Inches, RGBColor

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Pt(0)

    if fig_data.image_path and os.path.isfile(fig_data.image_path):
        run = p.add_run()
        run.add_picture(fig_data.image_path, width=Inches(5.0))
    else:
        # Placeholder box
        run = p.add_run(f'【图占位：图{fig_data.number} {fig_data.caption}】')
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string('999999')
        set_east_asian_font(run._element, 'Times New Roman', 'SimSun', 10)

    # Caption below the figure
    caption_text = f'图{fig_data.number} {fig_data.caption}'
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.first_line_indent = Pt(0)
    cp.paragraph_format.space_before = Pt(4)

    crun = cp.add_run(caption_text)
    crun.font.size = Pt(9)
    set_east_asian_font(crun._element, 'Times New Roman', 'SimSun', 9)
